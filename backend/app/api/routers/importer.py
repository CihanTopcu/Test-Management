"""Bulk case import from CSV or Excel.

Test cases arrive from everywhere -- a business analyst's spreadsheet, an
export from another tool, a colleague's CSV -- and TestRail's importer was one
of the things the team used most. The two halves of it are a preview that
guesses the column mapping and shows what would happen, and a commit that
applies exactly the mapping the user confirmed.
"""
import csv
import io
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ...db import get_session
from ...models import (Case, CaseStep, CaseType, CustomField, Priority,
                       Section, Suite, User)
from ..deps import current_user
from ..permissions import WRITE_CASES, assert_can

router = APIRouter(prefix="/api", tags=["import"])

MAX_ROWS = 5000

# The built-in targets a column can be mapped onto, and the header names we
# recognise for each without being told.
BUILTIN: dict[str, list[str]] = {
    "title": ["title", "baslik", "başlık", "case", "test case", "name", "ad",
              "summary", "ozet", "özet"],
    "section": ["section", "bolum", "bölüm", "folder", "klasor", "klasör",
                "suite section", "group", "grup"],
    "type": ["type", "tip", "case type", "test type"],
    "priority": ["priority", "oncelik", "öncelik"],
    "refs": ["refs", "references", "referans", "referanslar", "jira", "issue"],
    "estimate": ["estimate", "tahmin", "sure", "süre", "duration"],
    "steps": ["steps", "adim", "adım", "adimlar", "adımlar", "step",
              "test steps", "senaryo"],
    "expected": ["expected", "beklenen", "expected result", "beklenen sonuc",
                 "beklenen sonuç", "result"],
}


def _norm(value: str) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _read_table(raw: bytes, filename: str) -> tuple[list[str], list[list[str]]]:
    """Return (headers, rows) from a CSV or XLSX upload."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:  # pragma: no cover - depends on the deployment
            raise HTTPException(400, "Excel destegi kurulu degil, CSV yukleyin")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = book.active
        rows = [[("" if c is None else str(c)) for c in row]
                for row in sheet.iter_rows(values_only=True)]
        book.close()
    else:
        # Excel on a Turkish Windows writes UTF-8 with a BOM and semicolons;
        # sniffing the delimiter is what keeps those files from arriving as
        # one column per row.
        text = raw.decode("utf-8-sig", errors="replace")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]

    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise HTTPException(400, "dosya bos")
    headers = [str(h).strip() for h in rows[0]]
    return headers, rows[1:]


def _auto_map(session: Session, headers: list[str]) -> dict[str, str]:
    """Guess a column -> field mapping from the header row."""
    fields = session.scalars(
        select(CustomField).where(CustomField.entity == "case")).all()
    by_label = {_norm(f.label): f.system_name for f in fields}
    by_system = {_norm(f.system_name): f.system_name for f in fields}

    mapping: dict[str, str] = {}
    for header in headers:
        key = _norm(header)
        if not key:
            continue
        hit = next((target for target, names in BUILTIN.items()
                    if key in names), None)
        if hit is None:
            hit = by_label.get(key) or by_system.get(key) or by_system.get(
                _norm("custom_" + key))
        if hit:
            mapping[header] = hit
    return mapping


class ImportResult(BaseModel):
    headers: list[str]
    mapping: dict[str, str]
    total_rows: int
    would_create: int
    new_sections: list[str]
    skipped: list[str]
    sample: list[dict]
    created: int = 0


@router.get("/import/targets")
def import_targets(session: Session = Depends(get_session),
                   _: User = Depends(current_user)):
    """What a column can be mapped onto, for the mapping dropdowns."""
    fields = session.scalars(
        select(CustomField).where(CustomField.entity == "case")
        .order_by(CustomField.display_order)).all()
    return {
        "builtin": [{"key": k, "label": k} for k in BUILTIN],
        "custom": [{"key": f.system_name, "label": f.label,
                    "field_type": f.field_type} for f in fields],
    }


@router.post("/suites/{suite_id}/import", response_model=ImportResult)
async def import_cases(suite_id: int,
                       file: UploadFile = File(...),
                       mapping_json: str | None = Form(None),
                       section_id: int | None = Form(None),
                       create_sections: bool = Form(True),
                       dry_run: bool = Form(True),
                       session: Session = Depends(get_session),
                       user: User = Depends(current_user)):
    """Preview (dry_run) or apply a spreadsheet import."""
    import json

    suite = session.get(Suite, suite_id)
    if suite is None:
        raise HTTPException(404, "suite bulunamadi")
    assert_can(session, user, WRITE_CASES, suite.project_id)

    headers, rows = _read_table(await file.read(), file.filename or "")
    if len(rows) > MAX_ROWS:
        raise HTTPException(
            400, f"{len(rows)} satir cok fazla, en fazla {MAX_ROWS} satir")

    mapping = _auto_map(session, headers)
    if mapping_json:
        # the user's confirmed mapping wins over the guess, including the
        # columns they deliberately set to "ignore"
        mapping = {k: v for k, v in json.loads(mapping_json).items() if v}

    index = {h: i for i, h in enumerate(headers)}
    title_col = next((h for h, t in mapping.items() if t == "title"), None)
    if title_col is None:
        raise HTTPException(400, "basliga karsilik gelen bir sutun secilmeli")

    types = {_norm(t.name): t.id for t in session.scalars(select(CaseType))}
    priorities = {_norm(p.name): p.id for p in session.scalars(select(Priority))}
    sections = {_norm(s.name): s for s in session.scalars(
        select(Section).where(Section.suite_id == suite_id))}
    default_section = session.get(Section, section_id) if section_id else None
    if default_section is None:
        # A file with no section column and no section chosen would otherwise
        # skip every row. Land it somewhere obvious instead.
        default_section = session.scalar(
            select(Section).where(Section.suite_id == suite_id,
                                  Section.is_deleted.is_(False))
            .order_by(Section.display_order).limit(1))
        if default_section is None and not dry_run:
            default_section = Section(suite_id=suite_id, name="İçe aktarılan",
                                      depth=0, display_order=1)
            session.add(default_section)
            session.flush()

    def cell(row: list[str], header: str) -> str:
        i = index.get(header)
        return str(row[i]).strip() if i is not None and i < len(row) else ""

    new_sections: list[str] = []
    skipped: list[str] = []
    sample: list[dict] = []
    created = 0
    order = session.scalar(
        select(func.coalesce(func.max(Section.display_order), 0))
        .where(Section.suite_id == suite_id)) or 0
    now = datetime.now(timezone.utc)

    for number, row in enumerate(rows, start=2):
        title = cell(row, title_col)
        if not title:
            skipped.append(f"satir {number}: baslik bos")
            continue

        target = default_section
        section_name = next(
            (cell(row, h) for h, t in mapping.items() if t == "section"), "")
        if section_name:
            found = sections.get(_norm(section_name))
            if found is None:
                if not create_sections:
                    skipped.append(f"satir {number}: '{section_name}' bolumu yok")
                    continue
                if section_name not in new_sections:
                    new_sections.append(section_name)
                if dry_run:
                    found = None
                else:
                    order += 1
                    found = Section(suite_id=suite_id, name=section_name,
                                    depth=0, display_order=order)
                    session.add(found)
                    session.flush()
                    sections[_norm(section_name)] = found
            target = found or target

        if target is None and not dry_run and not section_name:
            skipped.append(f"satir {number}: hedef bolum yok")
            continue

        custom: dict = {}
        steps: list[tuple[str, str]] = []
        type_id = priority_id = None
        refs = estimate = None
        step_text = expected_text = ""

        for header, target_key in mapping.items():
            value = cell(row, header)
            if not value:
                continue
            if target_key == "type":
                type_id = types.get(_norm(value))
            elif target_key == "priority":
                priority_id = priorities.get(_norm(value))
            elif target_key == "refs":
                refs = value[:1000]
            elif target_key == "estimate":
                estimate = value[:50]
            elif target_key == "steps":
                step_text = value
            elif target_key == "expected":
                expected_text = value
            elif target_key in ("title", "section"):
                continue
            else:
                custom[target_key] = value

        # one cell holding several numbered lines is how spreadsheets carry
        # steps; split it so they land as real steps rather than a blob
        if step_text:
            lines = [ln.strip() for ln in step_text.splitlines() if ln.strip()]
            expected_lines = [ln.strip() for ln in expected_text.splitlines()
                              if ln.strip()]
            if len(lines) <= 1:
                steps = [(step_text, expected_text)]
            else:
                steps = [(ln, expected_lines[i] if i < len(expected_lines) else "")
                         for i, ln in enumerate(lines)]
        elif expected_text:
            steps = [("", expected_text)]

        if len(sample) < 5:
            sample.append({
                "title": title, "section": section_name or (
                    target.name if target else ""),
                "type_id": type_id, "priority_id": priority_id,
                "steps": len(steps), "custom": custom,
            })

        if dry_run:
            created += 1
            continue

        case = Case(
            suite_id=suite_id, section_id=target.id, title=title[:1000],
            type_id=type_id, priority_id=priority_id, refs=refs,
            estimate=estimate, custom=custom, is_deleted=False,
            created_by=user.id, created_on=now,
            updated_by=user.id, updated_on=now,
        )
        session.add(case)
        session.flush()
        for idx, (content, expected) in enumerate(steps):
            session.add(CaseStep(case_id=case.id, idx=idx, content=content,
                                 expected=expected))
        created += 1

    if dry_run:
        session.rollback()
    else:
        session.commit()

    return ImportResult(
        headers=headers, mapping=mapping, total_rows=len(rows),
        would_create=created if dry_run else 0,
        new_sections=new_sections, skipped=skipped[:50], sample=sample,
        created=0 if dry_run else created,
    )
